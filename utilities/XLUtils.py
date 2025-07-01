import openpyxl

def getRowCount(file, sheetname):
    """
    Returns the number of rows in a given sheet of an Excel file.

    :param file: Path to the Excel file.
    :param sheetname: Name of the sheet to count rows in.
    :return: Number of rows in the specified sheet.
    """
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetname]
    return sheet.max_row
def getColumnCount(file, sheetname):
    """
    Returns the number of columns in a given sheet of an Excel file.

    :param file: Path to the Excel file.
    :param sheetname: Name of the sheet to count columns in.
    :return: Number of columns in the specified sheet.
    """
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetname]
    return sheet.max_column

def readData(file, sheetname, rowno, columnno):
    """
    Reads data from a specific cell in a given sheet of an Excel file.

    :param file: Path to the Excel file.
    :param sheetname: Name of the sheet to read data from.
    :param rowno: Row number (1-based index).
    :param columnno: Column number (1-based index).
    :return: Data from the specified cell.
    """
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetname]
    return sheet.cell(row=rowno, column=columnno).value

def writeData(file, sheetname, rowno, columnno, data):
    """
    Writes data to a specific cell in a given sheet of an Excel file.

    :param file: Path to the Excel file.
    :param sheetname: Name of the sheet to write data to.
    :param rowno: Row number (1-based index).
    :param columnno: Column number (1-based index).
    :param data: Data to write into the specified cell.
    """
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetname]
    sheet.cell(row=rowno, column=columnno).value = data
    workbook.save(file)  # Save changes to the file

def getCellData(file, sheetname, rowno, columnno):
    """
    Gets the data from a specific cell in a given sheet of an Excel file.

    :param file: Path to the Excel file.
    :param sheetname: Name of the sheet to read data from.
    :param rowno: Row number (1-based index).
    :param columnno: Column number (1-based index).
    :return: Data from the specified cell.
    """
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetname]
    return sheet.cell(row=rowno, column=columnno).value